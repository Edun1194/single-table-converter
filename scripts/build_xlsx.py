#!/usr/bin/env python3
"""build_xlsx.py — 把整理好的 row list 写入标准模板 .xlsx.

输入: --input rows.json (list[dict])
模板: --template assets/templates/template1-{cn,en}.xlsx
输出: --output out.xlsx

模板字段顺序(1-idx) — 与 app 端 (excel-app-code) 现行模板一致:
  1  品牌 / Brand
  2  图片 / Image
  3  货号 / International code
  4  描述 / Described
  5  类目 / Category
  6  材质 / Material
  7  性别 / Gender
  8  颜色 / Color
  9  色号 / Color Code
  10 季节 / Season
  11 产地 / Made in
  12 零售价 / Retail price
  13 cost / cost
  14 C+- / C+-
  15 批发价 / Wholesale price    (= cost × (1 + C+-/100))
  16 尺码 / Size
  17 数量 / Quantity
  18 云仓 / Vendor ID

与 app 端保持一致的 4 条硬规则(2026-10-29 对齐,详见 references/app-parity.md):
  1. 图片显示高度 = 行高(默认 92pt),等比缩放,**原图字节不动**(清晰度保留)
  2. C+- 写**纯整数**(18 / -18),不带 + 号、不带 % 、不是字符串
  3. 批发价 = cost × (1 + C+-/100);**没有 C+ 就留空,不填 0**
  4. 云仓取 vendor_id(与 app 端 row 字段同名),兼容旧键 warehouse_id

用法:
  python3 build_xlsx.py --input rows.json --template template1-cn.xlsx --output out.xlsx
  python3 build_xlsx.py --input rows.json --language en --output out.xlsx
"""
from __future__ import annotations

import argparse
import json
import re
import shutil
import sys
import tempfile
from pathlib import Path
from typing import Dict, List, Optional, Sequence

import openpyxl
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.utils import get_column_letter

# 模板字段顺序(1-idx) - 与 SKILL.md 严格对齐
TEMPLATE_COLS_CN = [
    "品牌", "图片", "货号", "描述", "类目", "材质", "性别", "颜色", "色号",
    "季节", "产地", "零售价", "cost", "C+-", "批发价",
    "尺码", "数量", "云仓",
]
TEMPLATE_COLS_EN = [
    "Brand", "Image", "International code", "Described", "Category",
    "Material", "Gender", "Color", "Color Code", "Season", "Made in",
    "Retail price", "cost", "C+-", "Wholesale price",
    "Size", "Quantity", "Vendor ID",
]

# 18 列模板固定宽(列 2 图片宽=40)
COL_WIDTHS = {
    1: 16, 2: 40, 3: 26, 4: 38, 5: 14, 6: 18, 7: 12, 8: 14, 9: 14,
    10: 12, 11: 12, 12: 14, 13: 12, 14: 8, 15: 14, 16: 8, 17: 8, 18: 12,
}

# ---------------------------------------------------------------------------
# 单位换算(app 端同款)
# ---------------------------------------------------------------------------
EMU_PER_PX = 9525    # 1 px @96DPI = 9525 EMU
EMU_PER_PT = 12700   # 1 pt = 12700 EMU
IMAGE_GAP_PX = 8     # 同一行多图之间的间距


# ---------------------------------------------------------------------------
# 数值清洗
# ---------------------------------------------------------------------------

def _to_str_or_empty(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def _to_float_or_none(v) -> Optional[float]:
    if v is None or v == "":
        return None
    if isinstance(v, str):
        v = v.strip().replace(",", ".") if re.fullmatch(r"-?[\d.]+,\d+", v.strip()) else v.strip()
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def _to_int_or_none(v) -> Optional[int]:
    if v is None or v == "":
        return None
    try:
        return int(round(float(v)))
    except (TypeError, ValueError):
        return None


_CPLUS_RE = re.compile(r"^[+-]?\d+(?:\.\d+)?$")


def parse_cplus_int(v) -> Optional[int]:
    """把 C+ 各种写法归一成**纯整数**:18 / -18 / '18' / 'C+18%' / 'C-15%' / 0.18。

    - 'C+18%' / 'c+18'  → 18
    - 'C-15%' / '-15'   → -15
    - '18'              → 18
    - 0.18              → 18(小数比例)
    - '' / None / 非法   → None(留空)

    ⚠️ 返回 None 表示**没有 C+**,调用方必须留空,不能填 0。
    """
    if v is None:
        return None
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        f = float(v)
        # 0.18 这种小数比例 → 18;已经是 18 这种整数 → 18
        if 0 < abs(f) < 1:
            return int(round(f * 100))
        return int(round(f))
    s = str(v).strip()
    if not s:
        return None
    s = s.replace("％", "%")
    neg = False
    m = re.search(r"C\s*([+-])\s*(\d+(?:\.\d+)?)\s*%?", s, re.IGNORECASE)
    if m:
        neg = m.group(1) == "-"
        num = float(m.group(2))
    elif _CPLUS_RE.match(s):
        num = float(s)
        neg = num < 0
    else:
        return None
    if neg:
        num = -abs(num)
    if 0 < abs(num) < 1:
        num *= 100
    return int(round(num))


def compute_supply_price(
    cost: Optional[float], cplus_int: Optional[int]
) -> Optional[float]:
    """批发价 = cost × (1 + C+/100)。**没有 C+ 或没有 cost 就留空(None),不填 0。**"""
    if cost is None or cplus_int is None:
        return None
    return round(cost * (1 + cplus_int / 100.0), 2)


# ---------------------------------------------------------------------------
# 图片(与 app 端 embed_images_in_row 同款)
# ---------------------------------------------------------------------------

def _open_image(path: Path, tmp_dir: Path) -> Optional[XLImage]:
    """打开图片文件。openpyxl 不支持的格式(如 jp2)先转 PNG。"""
    try:
        return XLImage(str(path))
    except Exception:
        pass
    try:
        from PIL import Image as PILImage

        src = PILImage.open(str(path))
        if src.mode in ("RGBA", "LA", "P"):
            src = src.convert("RGB")
        out = tmp_dir / (path.stem + ".png")
        src.save(out, format="PNG")
        return XLImage(str(out))
    except Exception as exc:
        print(f"⚠️  image open failed for {path}: {exc}", file=sys.stderr)
        return None


def embed_images(
    ws,
    row_1idx: int,
    col_1idx: int,
    image_paths: Sequence[Path],
    max_height_emu: Optional[int],
    tmp_dir: Path,
) -> None:
    """把图片写进指定单元格。

    - 显示尺寸 = 原图像素 × 9525 EMU,**等比缩到不超过行高**(app 端同款)
    - **原图字节一字不动** —— 只改 openpyxl 写入的 ext,Excel 内部高质量缩放,清晰度保留
    - 同一行多图:按显示宽度累加 colOff 偏移,图之间留 8px 间距
    """
    running_col_off_emu = 0
    for path in image_paths:
        img = _open_image(path, tmp_dir)
        if img is None:
            continue
        # ⚠️ openpyxl 的 Image.width/height 是**像素**,不是 EMU。
        #    写 img.width = X * 9525 会双倍乘(图会变成 79ft 高)——绝不能那样写。
        cx = int(img.width * EMU_PER_PX)
        cy = int(img.height * EMU_PER_PX)
        if max_height_emu and cy > max_height_emu:
            scale = max_height_emu / cy
            cx = int(cx * scale)
            cy = max_height_emu
        img.anchor = OneCellAnchor(
            _from=AnchorMarker(
                col=col_1idx - 1,
                row=row_1idx - 1,
                colOff=running_col_off_emu,
                rowOff=0,
            ),
            ext=XDRPositiveSize2D(cx=cx, cy=cy),
        )
        ws.add_image(img)
        # 用**显示宽度** cx 累加偏移,不能用 img.width(原图宽度),
        # 否则第二张图会被推到很远
        running_col_off_emu += cx + int(IMAGE_GAP_PX * EMU_PER_PX)


def _collect_image_paths(fields: Dict, image_dir: Optional[Path]) -> List[Path]:
    """收图:支持 image_path(单图)和 image_paths(多图列表)。"""
    raw: List[str] = []
    if fields.get("image_path"):
        raw.append(str(fields["image_path"]))
    for p in fields.get("image_paths") or []:
        raw.append(str(p))
    out: List[Path] = []
    for s in raw:
        p = Path(s)
        if not p.is_absolute() and image_dir:
            p = image_dir / p
        if p.exists():
            out.append(p)
        else:
            print(f"⚠️  image not found: {p}", file=sys.stderr)
    return out


# ---------------------------------------------------------------------------
# 写行
# ---------------------------------------------------------------------------

def _write_row(
    ws,
    row_1idx: int,
    fields: Dict,
    headers: List[str],
    image_dir: Optional[Path],
    image_row_height: Optional[float],
    tmp_dir: Path,
) -> None:
    """把 fields 按 headers 顺序写到 row_1idx 行(图片单独 embed)。"""
    # ---- 先算好派生字段 ----
    cost = _to_float_or_none(fields.get("cost"))
    retail = _to_float_or_none(fields.get("retail_price"))
    cplus_int = parse_cplus_int(fields.get("cplus_text"))

    # 批发价:优先用调用方给的;没给但有 cost + C+ 就自己算;都没有 → 留空
    supply = _to_float_or_none(fields.get("supply_price"))
    if supply is None and cost is not None and cplus_int is not None:
        supply = compute_supply_price(cost, cplus_int)

    # 云仓:app 端 row 字段叫 vendor_id,兼容旧键 warehouse_id
    warehouse = fields.get("vendor_id")
    if warehouse in (None, ""):
        warehouse = fields.get("warehouse_id", "")

    image_paths = _collect_image_paths(fields, image_dir)

    for c_1idx, header in enumerate(headers, start=1):
        if header in ("Image", "图片"):
            continue  # 图片单独处理
        if header in ("International code", "货号"):
            val = _to_str_or_empty(fields.get("sku", ""))
        elif header in ("Brand", "品牌"):
            val = _to_str_or_empty(fields.get("brand", ""))
        elif header in ("Described", "描述"):
            val = _to_str_or_empty(fields.get("description", ""))
        elif header in ("Category", "类目"):
            val = _to_str_or_empty(fields.get("category", ""))
        elif header in ("Material", "材质"):
            val = _to_str_or_empty(fields.get("material", ""))
        elif header in ("Color", "颜色"):
            val = _to_str_or_empty(fields.get("color", ""))
        elif header in ("Color Code", "色号"):
            val = _to_str_or_empty(fields.get("color_code", ""))
        elif header in ("Gender", "性别"):
            val = _to_str_or_empty(fields.get("gender", ""))
        elif header in ("Season", "季节"):
            val = _to_str_or_empty(fields.get("season", ""))
        elif header in ("Made in", "产地"):
            val = _to_str_or_empty(fields.get("origin", ""))
        elif header == "cost":
            val = cost if cost is not None else ""
        elif header == "C+-":
            # 纯整数(模板规则):18 / -18;没有就留空
            val = cplus_int if cplus_int is not None else ""
        elif header in ("Wholesale price", "批发价"):
            val = supply if supply is not None else ""
        elif header in ("Retail price", "零售价"):
            val = retail if retail is not None else ""
        elif header in ("Size", "尺码"):
            val = _to_str_or_empty(fields.get("size", ""))
        elif header in ("Quantity", "数量"):
            q = _to_int_or_none(fields.get("quantity"))
            val = q if q is not None else ""
        elif header in ("Vendor ID", "云仓"):
            val = _to_str_or_empty(warehouse)
        else:
            val = _to_str_or_empty(fields.get(key_fallback(header), ""))
        ws.cell(row_1idx, c_1idx, val)

    # ---- 图片(有图才设行高,无图沿用模板默认)----
    if image_paths:
        max_height_emu = (
            int(image_row_height * EMU_PER_PT) if image_row_height else None
        )
        embed_images(ws, row_1idx, 2, image_paths, max_height_emu, tmp_dir)
        if image_row_height:
            ws.row_dimensions[row_1idx].height = float(image_row_height)


def key_fallback(header: str) -> str:
    return header.lower().replace(" ", "_")


# ---------------------------------------------------------------------------
# 主流程
# ---------------------------------------------------------------------------

def build(
    input_file: Path,
    template_file: Optional[Path],
    output_file: Path,
    language: str = "cn",
    image_row_height: Optional[float] = 92.0,
    image_col_width: Optional[float] = None,
    keep_zero_qty: bool = False,
) -> int:
    # 1) 读 rows.json
    raw = json.loads(input_file.read_text(encoding="utf-8"))
    if not isinstance(raw, list):
        raise ValueError("rows.json must be a list of dicts")

    # 2) 选模板
    if template_file is None:
        skill_root = Path(__file__).resolve().parent.parent
        template_file = skill_root / "assets" / "templates" / f"template1-{language}.xlsx"
    if not template_file.exists():
        raise FileNotFoundError(f"template not found: {template_file}")

    headers = TEMPLATE_COLS_CN if language == "cn" else TEMPLATE_COLS_EN

    # 3) load template
    wb = openpyxl.load_workbook(str(template_file))
    ws = wb.active

    # 4) 设置列宽
    for c_1idx, w in COL_WIDTHS.items():
        ws.column_dimensions[get_column_letter(c_1idx)].width = w
    if image_col_width:
        ws.column_dimensions[get_column_letter(2)].width = float(image_col_width)

    # 5) 写 rows (从第 2 行开始,R1 是 header)
    #    ⚠️ 新模板 max_row=140 是**空样式占位**,不能拿 max_row 当起点(会写到 R141)
    start_row = 2
    image_dir = input_file.parent
    tmp_dir = Path(tempfile.mkdtemp(prefix="stc_"))
    written = 0
    skipped = 0
    try:
        for i, fields in enumerate(raw):
            # qty=0 默认跳过(app 端 is_positive_qty 语义);
            # quantity 为 None =「还没填」(期货单),保留该行
            qty = _to_int_or_none(fields.get("quantity"))
            if not keep_zero_qty and qty is not None and qty <= 0:
                skipped += 1
                continue
            _write_row(
                ws, start_row + written, fields, headers, image_dir,
                image_row_height, tmp_dir,
            )
            written += 1
    finally:
        shutil.rmtree(tmp_dir, ignore_errors=True)

    # 6) 写文件
    output_file.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(output_file))
    if skipped:
        print(f"ℹ️  skipped {skipped} row(s) with quantity <= 0 (use --keep-zero-qty to keep)")
    return written


def main() -> None:
    p = argparse.ArgumentParser(description="Build single-table-converter output xlsx")
    p.add_argument("--input", required=True, help="rows.json path")
    p.add_argument("--template", help="target template .xlsx path")
    p.add_argument("--language", choices=["cn", "en"], default="cn", help="template language if --template not given")
    p.add_argument("--output", required=True, help="output .xlsx path")
    p.add_argument("--image-row-height", type=float, default=92.0,
                   help="有图数据行的行高(pt),图片等比缩到不超这个高度。默认 92")
    p.add_argument("--image-col-width", type=float, default=None,
                   help="图片列列宽(Excel 字符单位)。默认用模板/内置值 40")
    p.add_argument("--keep-zero-qty", action="store_true",
                   help="保留数量=0 的行(默认跳过,与 app 端 is_positive_qty 一致)")
    args = p.parse_args()

    n = build(
        Path(args.input),
        Path(args.template) if args.template else None,
        Path(args.output),
        language=args.language,
        image_row_height=args.image_row_height,
        image_col_width=args.image_col_width,
        keep_zero_qty=args.keep_zero_qty,
    )
    print(f"✅ wrote {n} rows → {args.output}")


if __name__ == "__main__":
    main()
